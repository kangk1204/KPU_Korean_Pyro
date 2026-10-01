import json
import unittest
from pathlib import Path

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]


class TablesContractTest(unittest.TestCase):
    def test_tables_json_schema_and_main_rows(self):
        payload = json.loads((ROOT / "tables/tables.json").read_text())
        self.assertEqual(list(payload.keys()), ["tables"])
        self.assertEqual([t["id"] for t in payload["tables"]], ["Table 1", "Table 2", "Table 3", "Table 4"])
        expected_rows = {"Table 1": 28, "Table 2": 10, "Table 3": 5, "Table 4": 10}
        for table in payload["tables"]:
            self.assertEqual(set(table.keys()), {"id", "title", "headers", "rows", "footnote"})
            self.assertEqual(len(table["rows"]), expected_rows[table["id"]])
            self.assertTrue(all(len(row) == len(table["headers"]) for row in table["rows"]))
        table1 = payload["tables"][0]
        self.assertIn("All patients (n=87)", table1["headers"])
        self.assertIn("Primary recurrence-analysis subset (n=82)", table1["headers"])
        self.assertIn("at least 3 cycles", table1["footnote"])
        self.assertIn("CA19-9 is reported in U/mL", table1["footnote"])
        table3 = payload["tables"][2]
        self.assertIn("Platform", table3["headers"])
        self.assertIn("Unique probes available", table3["headers"])
        self.assertIn("EPIC GPL21145", {row[1] for row in table3["rows"]})
        self.assertIn("450K GPL13534", {row[1] for row in table3["rows"]})
        self.assertIn("TCGA frozen feature snapshot", {row[1] for row in table3["rows"]})
        self.assertIn("77/77", {row[7] for row in table3["rows"]})
        self.assertIn("71/77", {row[7] for row in table3["rows"]})
        self.assertIn("not reverified", {row[7] for row in table3["rows"]})
        self.assertIn(">=80% valid-probe scoring rule is verified for these two independent paired public cohorts", table3["footnote"])
        self.assertIn("not independent of TCGA COADREAD", table3["footnote"])

    def test_workbooks_exist_and_protect_private_patient_identifiers(self):
        for rel in ["supplement/SupplementaryTables.xlsx", "supplement/SourceData.xlsx"]:
            path = ROOT / rel
            self.assertTrue(path.exists(), rel)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            self.assertIn("README", wb.sheetnames)
            readme = "\n".join(str(wb["README"].cell(i, 1).value or "") for i in range(1, 8))
            self.assertIn("Public release", readme)
        forbidden = {"patient_id", "operation_date", "endpoint_date", "left_out_patient_id"}
        source = pd.ExcelFile(ROOT / "supplement/SourceData.xlsx")
        for sheet in source.sheet_names:
            if sheet == "README":
                continue
            cols = set(pd.read_excel(source, sheet_name=sheet, nrows=0).columns)
            self.assertTrue(forbidden.isdisjoint(cols), sheet)
        supplementary = pd.ExcelFile(ROOT / "supplement/SupplementaryTables.xlsx")
        for sheet in ["clinical", "methylation_wide", "methylation_long", "paired_loo", "cutoff_loo", "cutoff_patient_calls"]:
            cols = set(pd.read_excel(supplementary, sheet_name=sheet, nrows=0).columns)
            self.assertTrue(forbidden.isdisjoint(cols), sheet)

    def test_main_csv_and_workbook_values_match_json(self):
        payload = json.loads((ROOT / "tables/tables.json").read_text())
        for table in payload["tables"]:
            rel = f"tables/{table['id'].lower().replace(' ', '_')}.csv"
            csv_df = pd.read_csv(ROOT / rel, dtype=str, keep_default_na=False)
            self.assertEqual(list(csv_df.columns), table["headers"])
            self.assertEqual(csv_df.values.tolist(), [[str(x) for x in row] for row in table["rows"]])
        wb = openpyxl.load_workbook(ROOT / "supplement/SupplementaryTables.xlsx", read_only=True, data_only=True)
        self.assertIn("cutoff_bootstrap", wb.sheetnames)
        self.assertIn("cutoff_loo", wb.sheetnames)
        self.assertIn("cutoff_cv", wb.sheetnames)
        self.assertEqual(wb["cutoff_bootstrap"].max_row, 61)
        self.assertEqual(wb["cutoff_loo"].max_row, 5221)
        self.assertEqual(wb["cutoff_cv"].max_row, 6001)


if __name__ == "__main__":
    unittest.main()
