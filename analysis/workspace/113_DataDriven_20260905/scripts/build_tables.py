"""Build manuscript tables and supplementary workbooks from frozen results."""
from __future__ import annotations

import csv
import json
import math
import platform
import re
import sys
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from common import GENES, ROOT, ensure_dirs, write_json



def require(condition, message="Integrity check failed"):
    if not condition:
        raise RuntimeError(message)

TABLE_DIR = ROOT / "tables"
SUPP_DIR = ROOT / "supplement"
RESULTS = ROOT / "results"
EXTERNAL = RESULTS / "external"


def read_table(path: Path) -> pd.DataFrame:
    sep = "\t" if path.suffix.lower() in {".tsv", ".txt"} else ","
    return pd.read_csv(path, sep=sep)


def fmt_num(x: Any, digits: int = 1) -> str:
    if pd.isna(x):
        return "NA"
    return f"{float(x):.{digits}f}"


def fmt_p(x: Any) -> str:
    if pd.isna(x):
        return "NA"
    x = float(x)
    if x < 0.001:
        return f"{x:.2e}"
    return f"{x:.3f}"


def fmt_n_pct(n: int, d: int) -> str:
    return f"{int(n)} ({100 * int(n) / int(d):.1f}%)"


def fmt_mean_sd(values: pd.Series, digits: int = 1) -> str:
    v = pd.to_numeric(values, errors="coerce").dropna()
    return f"{v.mean():.{digits}f} +/- {v.std(ddof=1):.{digits}f}"


def fmt_median_iqr(values: pd.Series, digits: int = 1) -> str:
    v = pd.to_numeric(values, errors="coerce").dropna()
    q1, q3 = np.quantile(v, [0.25, 0.75], method="linear")
    return f"{np.median(v):.{digits}f} ({q1:.{digits}f}-{q3:.{digits}f})"


def category_rows(df_all: pd.DataFrame, df_primary: pd.DataFrame, column: str, label: str) -> list[list[str]]:
    levels = sorted(set(df_all[column].dropna().astype(str)) | set(df_primary[column].dropna().astype(str)))
    rows = []
    for level in levels:
        rows.append(
            [
                f"{label}: {level}",
                fmt_n_pct((df_all[column].astype(str) == level).sum(), len(df_all)),
                fmt_n_pct((df_primary[column].astype(str) == level).sum(), len(df_primary)),
            ]
        )
    return rows


def build_table1(clinical: pd.DataFrame) -> dict[str, Any]:
    primary = clinical[clinical["recurrence_primary"].eq(1)].copy()
    rows: list[list[str]] = [
        ["Patients, n", str(len(clinical)), str(len(primary))],
        ["Age, mean +/- SD, years", fmt_mean_sd(clinical["age"]), fmt_mean_sd(primary["age"])],
        ["Age, median (IQR), years", fmt_median_iqr(clinical["age"]), fmt_median_iqr(primary["age"])],
    ]
    rows += category_rows(clinical, primary, "sex", "Sex")
    rows += category_rows(clinical, primary, "stage", "Provider stage")
    rows += [
        ["Lymphovascular invasion, n (%)", fmt_n_pct(clinical["lvi"].sum(), len(clinical)), fmt_n_pct(primary["lvi"].sum(), len(primary))],
        ["CEA, median (IQR), ng/mL", fmt_median_iqr(clinical["cea_ng_ml"]), fmt_median_iqr(primary["cea_ng_ml"])],
        ["CEA >7 ng/mL, n (%)", fmt_n_pct(clinical["cea_elevated"].sum(), len(clinical)), fmt_n_pct(primary["cea_elevated"].sum(), len(primary))],
        ["CA19-9, median (IQR), U/mL", fmt_median_iqr(clinical["ca199_u_ml"]), fmt_median_iqr(primary["ca199_u_ml"])],
        [
            "CA19-9 >37 U/mL, n (%)",
            fmt_n_pct(clinical["ca199_elevated"].sum(), len(clinical)),
            fmt_n_pct(primary["ca199_elevated"].sum(), len(primary)),
        ],
        ["PNI recorded positive, n (%)", fmt_n_pct(clinical["pni"].sum(), len(clinical)), fmt_n_pct(primary["pni"].sum(), len(primary))],
        [
            "Preoperative treatment recorded, n (%)",
            fmt_n_pct(clinical["preop_treatment"].sum(), len(clinical)),
            fmt_n_pct(primary["preop_treatment"].sum(), len(primary)),
        ],
        [
            "Postoperative chemotherapy code 1, n (%)",
            fmt_n_pct(clinical["postop_chemo"].sum(), len(clinical)),
            fmt_n_pct(primary["postop_chemo"].sum(), len(primary)),
        ],
        [
            "Recorded recurrence/progression, n (%)",
            fmt_n_pct(clinical["event"].sum(), len(clinical)),
            fmt_n_pct(primary["event"].sum(), len(primary)),
        ],
        ["Observed endpoint time, median (IQR), days", fmt_median_iqr(clinical["duration_days"], 0), fmt_median_iqr(primary["duration_days"], 0)],
    ]
    rows += category_rows(clinical, primary, "surgery", "Surgery")
    return {
        "id": "Table 1",
        "title": "Clinical characteristics of the pyrosequencing cohort",
        "headers": ["Characteristic", "All patients (n=87)", "Primary recurrence-analysis subset (n=82)"],
        "rows": rows,
        "footnote": (
            "Values are n (%), mean +/- SD, or median (IQR). The primary recurrence-analysis subset includes provider stage I-III "
            "patients without palliative surgery. Postoperative chemotherapy code 1 means at least 3 cycles; code 0 means fewer "
            "than 3 cycles or unknown, not confirmed absence of treatment. CA19-9 is reported in U/mL."
        ),
    }


def build_table2(paired: pd.DataFrame) -> dict[str, Any]:
    paired = paired.set_index("gene").loc[GENES].reset_index()
    rows = []
    for r in paired.itertuples(index=False):
        rows.append(
            [
                r.gene,
                f"{fmt_num(r.tumor_mean_pct)} +/- {fmt_num(r.tumor_sd_pct)}",
                f"{fmt_num(r.normal_mean_pct)} +/- {fmt_num(r.normal_sd_pct)}",
                fmt_num(r.mean_difference_pp),
                f"{fmt_num(r.mean_diff_bca95_low)} to {fmt_num(r.mean_diff_bca95_high)}",
                fmt_p(r.paired_t_BH_q),
                f"{int(r.tumor_greater_own_normal_n)}/{int(r.n_pairs)}",
            ]
        )
    return {
        "id": "Table 2",
        "title": "Paired tumor-normal pyrosequencing methylation differences",
        "headers": [
            "Gene",
            "Tumor methylation, mean +/- SD (%)",
            "Normal methylation, mean +/- SD (%)",
            "Mean difference, pp",
            "BCa 95% CI, pp",
            "Paired t BH q",
            "Tumor > own normal",
        ],
        "rows": rows,
        "footnote": (
            "Difference is tumor minus paired adjacent-normal methylation in percentage points. Confidence intervals are "
            "patient-paired BCa bootstrap intervals with 5000 replicates and a shared patient-index bootstrap matrix. "
            "BH q values correct the 10 paired t tests."
        ),
    }


def build_table3(effects: pd.DataFrame, coverage: pd.DataFrame, cohort_summary: pd.DataFrame) -> dict[str, Any]:
    platforms = {
        "colonomics": "450K GPL13534",
        "gse119526": "EPIC GPL21145",
        "tcga_coadread": "TCGA frozen feature snapshot",
        "tcga_coad": "TCGA frozen feature snapshot",
        "tcga_read": "TCGA frozen feature snapshot",
    }
    evidence_roles = {
        "colonomics": "Independent paired public cohort",
        "gse119526": "Independent paired public cohort",
        "tcga_coadread": "Supportive frozen-feature cohort; discovery overlap",
        "tcga_coad": "Nested TCGA subset; not independent",
        "tcga_read": "Nested TCGA subset; not independent",
    }
    cohort_order = ["colonomics", "gse119526", "tcga_coadread", "tcga_coad", "tcga_read"]
    cohort_labels = {
        "colonomics": "Colonomics", "gse119526": "GSE119526",
        "tcga_coadread": "TCGA COAD / READ", "tcga_coad": "TCGA COAD", "tcga_read": "TCGA READ",
    }
    summaries = cohort_summary.set_index("cohort")
    effect_summary = effects.groupby("cohort").agg(
        genes_estimable=("gene", "nunique"),
        n_pairs_min=("n_pairs", "min"),
        n_pairs_max=("n_pairs", "max"),
    )
    coverage_summary = coverage.groupby("cohort").agg(
        target_probes=("n_target_probes", "sum"),
        available_probes=("n_available_probes", "sum"),
        missing_available=("n_available_probes", lambda s: int(s.isna().sum())),
    )
    rows = []
    for cohort in cohort_order:
        c = summaries.loc[cohort]
        e = effect_summary.loc[cohort]
        cov = coverage_summary.loc[cohort]
        if int(cov["missing_available"]) > 0:
            availability = "not reverified"
        else:
            availability = f"{int(cov['available_probes'])}/{int(cov['target_probes'])}"
        pair_range = str(int(e["n_pairs_min"])) if int(e["n_pairs_min"]) == int(e["n_pairs_max"]) else f"{int(e['n_pairs_min'])}-{int(e['n_pairs_max'])}"
        rows.append(
            [
                cohort_labels[cohort],
                platforms.get(cohort, "NA"),
                evidence_roles.get(cohort, "NA"),
                str(int(c["n_patients"])),
                str(int(c["n_tn_pairs_by_patient"])),
                pair_range,
                f"{int(e['genes_estimable'])}/10",
                availability,
            ]
        )
    return {
        "id": "Table 3",
        "title": "External cohort overview and fixed probe coverage",
        "headers": [
            "Cohort",
            "Platform",
            "Evidence role",
            "Unique subject IDs in snapshot",
            "Paired tumor-normal patients",
            "Analyzed pairs per gene",
            "Genes estimable",
            "Unique probes available",
        ],
        "rows": rows,
        "footnote": (
            "Colonomics uses Infinium HumanMethylation450K (GPL13534); GSE119526 uses Infinium MethylationEPIC (GPL21145). "
            "Unique ID counts refer to all records in each snapshot, including unpaired specimens and, for Colonomics, healthy-mucosa records; paired analysis sizes are shown separately. "
            "The >=80% valid-probe scoring rule is verified for these two independent paired public cohorts. TCGA rows are "
            "supportive frozen-feature rows with discovery overlap; COAD and READ are nested subsets and not independent of "
            "TCGA COADREAD. TCGA probe availability is reported as not reverified. Detailed gene-level effects, minimum "
            "valid-probe counts, and probe lists are provided in SupplementaryTables.xlsx and SourceData.xlsx. Beta-scale "
            "effects are not pooled with pyrosequencing percentage-point effects."
        ),
    }


def build_table4(recurrence: pd.DataFrame) -> dict[str, Any]:
    primary = recurrence[recurrence["analysis_population"].eq("stage_I_III_nonpalliative_primary")]
    primary = primary.set_index("gene").loc[GENES].reset_index()
    rows = []
    for r in primary.itertuples(index=False):
        rows.append(
            [
                r.gene,
                str(int(r.complete_cases)),
                str(int(r.events)),
                fmt_num(r.hr_per10, 2),
                f"{fmt_num(r.ci_lower, 2)} to {fmt_num(r.ci_upper, 2)}",
                fmt_p(r.p_value),
                fmt_p(r.p_bh),
                fmt_p(r.ph_p_value),
                "Yes" if bool(r.estimable) else "No",
            ]
        )
    return {
        "id": "Table 4",
        "title": "Exploratory recurrence associations in the primary analysis subset",
        "headers": [
            "Gene",
            "Complete cases",
            "Events",
            "HR per 10 pp tumor methylation",
            "95% CI",
            "Cox p",
            "BH q",
            "PH test p",
            "Estimable",
        ],
        "rows": rows,
        "footnote": (
            "Unadjusted Cox proportional hazards models use provider stage I-III patients without palliative surgery "
            "(n=82, 14 events). Hazard ratios are per 10 percentage-point higher tumor methylation. Efron ties were used. "
            "These analyses are exploratory and were not used to choose cutoffs or genes."
        ),
    }


def write_csv(table: dict[str, Any]) -> None:
    path = TABLE_DIR / f"{table['id'].lower().replace(' ', '_')}.csv"
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(table["headers"])
        writer.writerows(table["rows"])


def write_md(table: dict[str, Any]) -> None:
    path = TABLE_DIR / f"{table['id'].lower().replace(' ', '_')}.md"
    lines = [f"# {table['id']}. {table['title']}", ""]
    headers = table["headers"]
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
    for row in table["rows"]:
        safe = [str(x).replace("|", "\\|") for x in row]
        lines.append("| " + " | ".join(safe) + " |")
    lines.extend(["", f"Note: {table['footnote']}", ""])
    path.write_text("\n".join(lines))


def clean_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\:\*\?\/\\]", "_", name)[:31] or "Sheet"
    candidate = base
    i = 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def write_df_sheet(wb: openpyxl.Workbook, name: str, df: pd.DataFrame, used: set[str]) -> None:
    ws = wb.create_sheet(clean_sheet_name(name, used))
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([None if pd.isna(x) else x for x in row])
    style_sheet(ws)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
            elif isinstance(cell.value, int):
                cell.number_format = "0"
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for idx, col in enumerate(ws.columns, start=1):
        width = min(max([len(str(c.value)) if c.value is not None else 0 for c in col] + [8]) + 2, 42)
        ws.column_dimensions[get_column_letter(idx)].width = width


def write_readme_sheet(wb: openpyxl.Workbook, used: set[str], title: str, lines: list[str]) -> None:
    ws = wb.active if wb.active.max_row == 1 and wb.active.max_column == 1 and wb.active["A1"].value is None else wb.create_sheet()
    ws.title = clean_sheet_name("README", used)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    for i, line in enumerate(lines, start=3):
        ws.cell(i, 1, line)
    ws.column_dimensions["A"].width = 110
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def table_to_df(table: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame(table["rows"], columns=table["headers"])


def public_patient_tables(clinical: pd.DataFrame, methylation_wide: pd.DataFrame, methylation_long: pd.DataFrame) -> dict[str, pd.DataFrame]:
    allowed_clinical = [
        "study_id",
        "age",
        "sex",
        "tnm",
        "stage",
        "lvi",
        "cea_ng_ml",
        "cea_elevated",
        "ca199_u_ml",
        "ca199_elevated",
        "pni",
        "preop_treatment",
        "surgery",
        "postop_chemo",
        "event",
        "duration_days",
        "recurrence_primary",
    ]
    public_wide_cols = ["study_id"] + [f"{t}_{g}" for g in GENES for t in ("T", "N")]
    return {
        "clinical_candidate": clinical[allowed_clinical].copy(),
        "methylation_wide_candidate": methylation_wide[public_wide_cols].copy(),
        "methylation_long_candidate": methylation_long[["study_id", "gene", "tissue", "methylation_pct"]].copy(),
    }


def scrub_patient_identifiers(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "patient_id" in out.columns:
        out = out.drop(columns=["patient_id"])
    for col in ["operation_date", "endpoint_date"]:
        if col in out.columns:
            out = out.drop(columns=[col])
    for col in ["left_out_patient_id"]:
        if col in out.columns:
            out = out.drop(columns=[col])
    return out


def write_workbooks(tables: list[dict[str, Any]], data: dict[str, pd.DataFrame]) -> None:
    supplementary = openpyxl.Workbook()
    used: set[str] = set()
    write_readme_sheet(
        supplementary,
        used,
        "Supplementary tables",
        [
            "Local full-result workbook for the 2026-09-05 data-driven CRC methylation reanalysis.",
            "Patient-level sheets use study_id and omit original patient_id, operation_date, and endpoint_date.",
            "Public release is pending author, IRB, and repository confirmation.",
            "Cutoff bootstrap confidence intervals are percentile tails, not BCa intervals.",
        ],
    )
    for table in tables:
        write_df_sheet(supplementary, table["id"].replace(" ", "_"), table_to_df(table), used)
    for name, df in data.items():
        write_df_sheet(supplementary, name, scrub_patient_identifiers(df), used)
    supplementary.save(SUPP_DIR / "SupplementaryTables.xlsx")

    source = openpyxl.Workbook()
    used = set()
    write_readme_sheet(
        source,
        used,
        "Source data",
        [
            "Candidate source-data workbook for table and figure reproduction.",
            "Public release pending: do not deposit externally until author, IRB, and repository checks are complete.",
            "Patient-level sheets use study_id only and exclude original patient IDs and actual dates.",
        ],
    )
    for table in tables:
        write_df_sheet(source, table["id"].replace(" ", "_"), table_to_df(table), used)
    for name, df in public_patient_tables(data["clinical"], data["methylation_wide"], data["methylation_long"]).items():
        write_df_sheet(source, name, df, used)
    for name in [
        "paired",
        "correlations",
        "pca_scores",
        "pca_loadings",
        "cutoffs",
        "cutoff_bootstrap",
        "recurrence",
        "recurrence_all_stage",
        "external_paired_effects",
        "external_probe_coverage",
        "external_cohort_summary",
    ]:
        if name in data:
            write_df_sheet(source, name, scrub_patient_identifiers(data[name]), used)
    source.save(SUPP_DIR / "SourceData.xlsx")


def workbook_summary(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {
        "path": str(path.relative_to(ROOT)),
        "sheets": wb.sheetnames,
        "sheet_count": len(wb.sheetnames),
        "max_dimensions": {ws.title: [ws.max_row, ws.max_column] for ws in wb.worksheets},
    }


def main() -> None:
    ensure_dirs()
    TABLE_DIR.mkdir(exist_ok=True)
    SUPP_DIR.mkdir(exist_ok=True)

    data = {
        "clinical": read_table(ROOT / "data/derived/clinical.tsv"),
        "methylation_wide": read_table(ROOT / "data/derived/methylation_wide.tsv"),
        "methylation_long": read_table(ROOT / "data/derived/methylation_long.tsv"),
        "legacy_calls": read_table(ROOT / "qc/legacy_calls.tsv"),
        "paired": read_table(RESULTS / "paired.csv"),
        "paired_loo": read_table(RESULTS / "paired_loo.csv"),
        "correlations": read_table(RESULTS / "correlations.csv"),
        "pca_scores": read_table(RESULTS / "pca_scores.csv"),
        "pca_loadings": read_table(RESULTS / "pca_loadings.csv"),
        "cutoffs": read_table(RESULTS / "cutoffs.csv"),
        "cutoff_bootstrap": read_table(RESULTS / "cutoff_bootstrap.csv"),
        "cutoff_loo": read_table(RESULTS / "cutoff_loo.csv"),
        "cutoff_cv": read_table(RESULTS / "cutoff_cv.csv"),
        "cutoff_patient_calls": read_table(RESULTS / "cutoff_patient_calls.tsv"),
        "recurrence": read_table(RESULTS / "recurrence.csv"),
        "recurrence_all_stage": read_table(RESULTS / "recurrence_all_stage.csv"),
        "survival_followup_summary": read_table(RESULTS / "survival_followup_summary.csv"),
        "survival_event_time_distribution": read_table(RESULTS / "survival_event_time_distribution.csv"),
        "survival_ph_diagnostics": read_table(RESULTS / "survival_ph_diagnostics.csv"),
        "external_cohort_summary": read_table(EXTERNAL / "cohort_summary.tsv"),
        "external_probe_coverage": read_table(EXTERNAL / "probe_coverage.tsv"),
        "external_paired_effects": read_table(EXTERNAL / "paired_public_effects.tsv"),
        "external_public_gene_scores": read_table(EXTERNAL / "public_gene_scores.tsv"),
    }
    require(len(data['clinical']) == 87, "Integrity check failed: len(data['clinical']) == 87")
    require(data['clinical']['recurrence_primary'].sum() == 82, "Integrity check failed: data['clinical']['recurrence_primary'].sum() == 82")
    require(data['clinical'].loc[data['clinical']['recurrence_primary'].eq(1), 'event'].sum() == 14, "Integrity check failed: data['clinical'].loc[data['clinical']['recurrence_primary'].eq(1), 'event'].sum() == 14")
    require(data['clinical']['cea_elevated'].sum() == 18, "Integrity check failed: data['clinical']['cea_elevated'].sum() == 18")

    tables = [
        build_table1(data["clinical"]),
        build_table2(data["paired"]),
        build_table3(data["external_paired_effects"], data["external_probe_coverage"], data["external_cohort_summary"]),
        build_table4(data["recurrence"]),
    ]
    for table in tables:
        write_csv(table)
        write_md(table)
    write_json(TABLE_DIR / "tables.json", {"tables": tables})
    write_workbooks(tables, data)

    summary = {
        "generated": {
            "tables_json": "tables/tables.json",
            "main_table_csv": [f"tables/{t['id'].lower().replace(' ', '_')}.csv" for t in tables],
            "main_table_md": [f"tables/{t['id'].lower().replace(' ', '_')}.md" for t in tables],
            "supplementary_tables": "supplement/SupplementaryTables.xlsx",
            "source_data": "supplement/SourceData.xlsx",
        },
        "table_row_counts": {t["id"]: len(t["rows"]) for t in tables},
        "workbooks": [
            workbook_summary(SUPP_DIR / "SupplementaryTables.xlsx"),
            workbook_summary(SUPP_DIR / "SourceData.xlsx"),
        ],
        "privacy": "patient-level public candidate sheets use study_id and exclude original patient_id, operation_date, endpoint_date",
        "versions": {
            "python": sys.version,
            "platform": platform.platform(),
            "pandas": pd.__version__,
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
        },
    }
    write_json(TABLE_DIR / "build_tables_summary.json", summary)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
