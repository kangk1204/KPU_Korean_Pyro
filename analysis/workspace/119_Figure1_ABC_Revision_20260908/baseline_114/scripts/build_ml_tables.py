"""Add two manuscript tables and a workbook from completed ML result files."""
from __future__ import annotations
import json
import math
import openpyxl
import pandas as pd
from common import ROOT, write_json
from build_tables import write_csv, write_md, write_df_sheet, table_to_df, write_readme_sheet

LABELS = {'ridge': 'Ridge', 'elastic_net': 'Elastic net', 'random_forest': 'Random forest',
          'survival_forest': 'Survival forest', 'rbf_svm': 'RBF SVM', 'best_single_gene': 'Inner-selected single gene'}
BLOCKS = {'clinical': 'Clinical', 'tumor10': 'Tumor methylation', 'methylation': 'Tumor methylation', 'combined': 'Combined'}


def number(value):
    return 'NE' if value is None or not math.isfinite(float(value)) else f'{float(value):.3f}'


def interval(rows, model, metric, *, block=None, horizon=None):
    candidates = [r for r in rows if r['model'] == model and r['metric'] == metric
                  and (block is None or r.get('block') == block)
                  and (horizon is None or r.get('horizon_days') == horizon)]
    if len(candidates) != 1:
        raise ValueError(f'Missing or duplicate CI: {model}, {metric}, {block}, {horizon}')
    r = candidates[0]
    return f"{number(r.get('point', r.get('estimate')))} ({number(r['ci_low'])} to {number(r['ci_high'])})"


def records_frame(records):
    return pd.DataFrame(records if records is not None else [])


def tissue_optimism_frame(summary):
    opt = summary['tissue'].get('optimism_corrected', {})
    rows = []
    mapping = {
        'auc': ('full_original_apparent_auc', 'mean_bootstrap_optimism_auc', 'optimism_corrected_auc'),
        'brier': ('full_original_apparent_brier', 'mean_bootstrap_optimism_brier', 'optimism_corrected_brier'),
    }
    for metric, (apparent, optimism, corrected) in mapping.items():
        rows.append({'model': opt.get('model', 'ridge'), 'metric': metric,
                     'original_apparent': opt.get(apparent),
                     'mean_optimism': opt.get(optimism),
                     'corrected': opt.get(corrected),
                     'B': opt.get('bootstrap_draws'),
                     'selection': opt.get('selection')})
    return pd.DataFrame(rows)


def recurrence_calibration_frame(recurrence):
    rows = []
    wanted = ['block', 'model', 'horizon_days', 'cal_estimable_repeats',
              'cal_intercept_mean', 'cal_slope_mean', 'ibs_status']
    for source, records in [('primary', recurrence['model_metrics']),
                            ('null_km', recurrence.get('null_km_metrics', [])),
                            ('sensitivity', recurrence.get('sensitivity_metrics', []))]:
        frame = records_frame(records)
        if frame.empty:
            continue
        keep = [c for c in wanted if c in frame.columns]
        out = frame[keep].copy()
        out.insert(0, 'source', source)
        rows.append(out)
    return pd.concat(rows, ignore_index=True) if rows else pd.DataFrame(columns=['source'] + wanted)


def maybe_read_frame(path):
    if not path.exists():
        return None
    return pd.read_csv(path, sep='\t' if '.tsv' in path.name else ',', keep_default_na='permutation' not in path.name)


def key_value_frame(mapping):
    rows = []
    for key, value in (mapping or {}).items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, sort_keys=True)
        rows.append({'key': key, 'value': value})
    return pd.DataFrame(rows)


def main():
    summary = json.loads((ROOT/'results/ml/summary.json').read_text())
    assert summary['status'] == 'final'
    tissue = summary['tissue']
    rows = []
    for cohort, n, result in [('Local PSQ', 87, tissue)] + [
            ('Colonomics', 92, summary['public']['colonomics']),
            ('GSE119526', 48, summary['public']['gse119526'])]:
        for r in result['model_metrics']:
            rows.append([cohort, LABELS[r['model']], str(n),
                         interval(result['bootstrap_ci'], r['model'], 'auc'), number(r['brier']),
                         number(r['threshold_0_5_sensitivity']), number(r['threshold_0_5_specificity'])])
    table5 = {'id': 'Table 5', 'title': 'Patient-grouped nested tissue classification within each cohort',
              'headers': ['Cohort', 'Model', 'Patient pairs', 'AUC (conditional 95% CI)',
                          'Brier score', 'Sensitivity at 0.5', 'Specificity at 0.5'], 'rows': rows,
              'footnote': 'Metrics are averaged across 20 repeats of patient-grouped nested cross-validation (outer 5 folds, inner 4 folds). '
                          'Intervals use 2000 patient-cluster resamples conditional on the fitted CV models. '
                          'The single-gene benchmark selects the gene within inner CV. Public-cohort models are fitted and evaluated within each cohort; '
                          'they do not validate a frozen pyrosequencing model across assays. NE denotes nonestimable or inapplicable; '
                          'the SVM provides ranking scores and therefore has no probability Brier score or 0.5 probability threshold.'}
    recurrence = summary['recurrence']
    rows = []
    for model in ['ridge', 'elastic_net', 'survival_forest']:
        for block in ['clinical', 'tumor10', 'combined']:
            matches = [r for r in recurrence['model_metrics'] if r['model'] == model and r['block'] in
                       ([block, 'methylation'] if block == 'tumor10' else [block]) and r['horizon_days'] == 1825]
            assert len(matches) == 1
            r = matches[0]
            rows.append([LABELS[model], BLOCKS[r['block']],
                         interval(recurrence['bootstrap_ci'], model, 'uno_c', block=r['block'], horizon=1825),
                         number(r.get('harrell_c_mean')), number(r.get('auc_mean')),
                         number(r.get('brier_mean')), number(r.get('ibs_mean'))])
    table6 = {'id': 'Table 6', 'title': 'Nested recurrence prediction in 82 patients with 14 recorded events',
              'headers': ['Model', 'Predictor block', '5-year Uno C (conditional 95% CI)', 'Harrell C',
                          '5-year AUC', '5-year Brier score', 'Integrated Brier score'], 'rows': rows,
              'footnote': 'Outer event-stratified 4-fold cross-validation was repeated 25 times with 3-fold inner tuning. '
                          'Clinical predictors were provider stage III, log1p CEA, and lymphovascular invasion. '
                          'The methylation block contained 10 continuous tumor measurements. '
                          'Censoring weights and survival probabilities were estimated using outer-training data only. '
                          'Concordance pools within-fold comparable-pair counts within each repeat; horizon scores pool cross-fitted probabilities within each repeat. '
                          'Intervals use 2000 patient resamples conditional on fitted models. Integrated Brier score covers 0–1825 days. '
                          'NE denotes nonestimability under the specified censoring-support rule. Death data were unavailable.'}
    payload = json.loads((ROOT/'tables/tables.json').read_text())
    tables = [t for t in payload['tables'] if t['id'] in ['Table 1','Table 2','Table 3','Table 4']] + [table5, table6]
    assert len(tables) == 6
    write_json(ROOT/'tables/tables.json', {'tables': tables})
    for table in [table5, table6]:
        write_csv(table)
        write_md(table)

    frames = {'Table_5': table_to_df(table5), 'Table_6': table_to_df(table6),
              'global_tests': pd.DataFrame([{'task': k, **v} for k, v in summary['global_tests'].items()]),
              'tissue_optimism_corrected': tissue_optimism_frame(summary),
              'recurrence_null_km_metrics': records_frame(recurrence.get('null_km_metrics')),
              'recurrence_optimism_corrected': records_frame(recurrence.get('optimism_corrected')),
              'recurrence_calibration': recurrence_calibration_frame(recurrence),
              'recurrence_regularization': records_frame(recurrence.get('regularization_diagnostics')),
              'recurrence_sensitivity_metrics': records_frame(recurrence.get('sensitivity_metrics')),
              'global_optimism': key_value_frame(summary.get('optimism')),
              'global_sensitivity': key_value_frame(summary.get('sensitivity'))}
    for prefix, folder, files in [
        ('tissue', ROOT/'results/ml/tissue', ['metrics.tsv','bootstrap_ci.tsv','ridge_permutation.tsv','ridge_optimism_bootstrap.tsv','oof_predictions.tsv.gz','folds.tsv']),
        ('recurrence', ROOT/'results/ml/recurrence', ['primary_metrics.csv','primary_bootstrap_ci.csv','primary_repeat_metrics.csv',
          'primary_bootstrap_repeat_mean_draws.csv','primary_null_km_metrics.csv','primary_ridge_permutation.csv',
          'primary_ridge_optimism_bootstrap.csv','primary_ridge_optimism_corrected.csv','ridge_sensitivity_metrics.csv',
          'primary_oof_predictions.csv','primary_folds.csv','primary_selected_hyperparameters.csv']),
        *[(name, ROOT/'results/ml/public'/name, ['metrics.tsv','bootstrap_ci.tsv','oof_predictions.tsv.gz','folds.tsv'])
          for name in ['colonomics','gse119526']]]:
        for filename in files:
            path = folder/filename
            frame = maybe_read_frame(path)
            if frame is None:
                continue
            if 'patient_id' in frame:
                raise ValueError(f'Export must use study_id column: {path}')
            assert not {'operation_date','endpoint_date','left_out_patient_id'} & set(frame.columns)
            if prefix == 'recurrence' and filename == 'primary_oof_predictions.csv':
                # The 202 IBS-grid columns remain in the full machine-readable CSV.
                # The review workbook keeps identifiers, folds and fixed-horizon predictions.
                frame = frame.loc[:, ~frame.columns.str.match(r'^(risk_grid_|G_grid_)')]
            frames[prefix+'_'+filename.split('.')[0]] = frame
        diag = maybe_read_frame(folder/'calibration_diagnostics.tsv')
        if diag is not None:
            frames[prefix+'_calibration_diagnostics'] = diag
    workbook = openpyxl.Workbook()
    used = set()
    write_readme_sheet(workbook, used, 'Machine learning results', [
        'Nested internal performance and within-public-cohort evaluation; not clinical validation.',
        'Public release pending author and institutional approval. Study identifiers only; no calendar dates.',
        'Bootstrap intervals are conditional on fitted cross-validation models. Repeated predictions are not independent patients.',
        'All planned model comparisons and sensitivity results are retained regardless of direction.',
        'Optimism-corrected rows are full-original apparent performance minus mean bootstrap optimism and are point-sensitivity summaries, not confidence intervals.',
        'Detailed model objects, training transformations and candidate tuning records remain in results/ml.',
        'The complete 101-point recurrence risk and censoring grids are in results/ml/recurrence/primary_oof_predictions.csv; the workbook OOF sheet retains fixed-horizon predictions.'])
    for name, frame in frames.items():
        write_df_sheet(workbook, name, frame, used)
    workbook.save(ROOT/'supplement/MLResults.xlsx')
    for filename in ['SupplementaryTables.xlsx', 'SourceData.xlsx']:
        path = ROOT/'supplement'/filename
        wb = openpyxl.load_workbook(path)
        for name in ['Table_5','Table_6','ML_global_tests']:
            if name in wb:
                del wb[name]
        used = set(wb.sheetnames)
        for table in [table5, table6]:
            write_df_sheet(wb, table['id'].replace(' ', '_'), table_to_df(table), used)
        write_df_sheet(wb, 'ML_global_tests', frames['global_tests'], used)
        wb.save(path)
    write_json(ROOT/'tables/ml_tables_summary.json', {'tables': ['Table 5','Table 6'],
               'row_counts': [len(table5['rows']), len(table6['rows'])], 'ml_workbook_sheets': list(workbook.sheetnames)})
    print('Added Tables 5/6 and MLResults.xlsx from completed canonical ML outputs.')


if __name__ == '__main__':
    main()
