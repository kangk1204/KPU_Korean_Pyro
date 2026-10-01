"""Verify frozen-source integrity, result contracts, and local delivery files."""
import json
from pathlib import Path
import subprocess
import sys
import zipfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from common import ROOT,GENES,N_BOOT,bh,sha256,write_json


def require(condition, message="Integrity check failed"):
    if not condition:
        raise RuntimeError(message)

STEMS=['F8_healthy_reference','F9_lesions_replication','F10_expression_context']


def check_baseline():
    manifest=json.loads((ROOT/'registry/baseline_manifest.json').read_text())
    failures=[]
    copies=[('snapshot',ROOT/'baseline_114')]
    if Path(manifest['source']).is_dir():copies.append(('original',Path(manifest['source'])))
    for entry in manifest['files']:
        for label,base in copies:
            path=base/entry['path']
            if not path.is_file() or path.stat().st_size!=entry['bytes'] or sha256(path)!=entry['sha256']:
                failures.append(dict(location=label,path=entry['path']))
    result=dict(files_per_copy=len(manifest['files']),failures=failures,original_present=Path(manifest['source']).exists())
    write_json(ROOT/'verification/baseline_integrity.json',result)
    if failures:raise AssertionError(f'Baseline integrity failure: {failures[:3]}')
    return result


def check_downloads():
    collections=[json.loads((ROOT/'registry/geo_downloads.json').read_text()),
        json.loads((ROOT/'registry/provenance_downloads.json').read_text())['downloads'],
        json.loads((ROOT/'registry/expression_manifest.json').read_text())['inputs']]
    checked=[]
    for records in collections:
        for r in records:
            p=ROOT/r['path']
            if not p.is_file() or sha256(p)!=r['sha256']:raise AssertionError(f'Source digest: {p}')
            checked.append(dict(path=r['path'],sha256=r['sha256'],bytes=p.stat().st_size))
    return checked


def check_results():
    run=json.loads((ROOT/'registry/biology_run.json').read_text())
    require(run['n_boot'] == N_BOOT and run['publication_run'], "Integrity check failed: run['n_boot'] == N_BOOT and run['publication_run']")
    d=pd.read_csv(ROOT/'results/tissue_contrasts.tsv',sep='\t')
    require(not d.duplicated(['cohort', 'gene', 'contrast']).any(), "Integrity check failed: not d.duplicated(['cohort', 'gene', 'contrast']).any()")
    require(d.groupby('family').size().to_dict() == {'healthy_reference': 60, 'lesion': 30, 'tissue_replication': 60}, "Integrity check failed: d.groupby('family').size().to_dict() == {'healthy_reference': 60, 'lesion': 30, 'tissue_replication': 60}")
    for _,g in d.groupby('family'):np.testing.assert_allclose(g.q_BH,bh(g.p),equal_nan=True,rtol=1e-12)
    estimated=d[d.status.eq('estimated')]
    require(estimated.bootstrap_replicates.eq(N_BOOT).all(), 'Integrity check failed: estimated.bootstrap_replicates.eq(N_BOOT).all()')
    require(np.isfinite(estimated[['effect', 'se', 'ci_low', 'ci_high', 'p', 'q_BH']].to_numpy()).all(), "Integrity check failed: np.isfinite(estimated[['effect', 'se', 'ci_low', 'ci_high', 'p', 'q_BH']].to_numpy()).all()")
    require((estimated.ci_low <= estimated.ci_high).all(), 'Integrity check failed: (estimated.ci_low <= estimated.ci_high).all()')
    for c in run['cohorts']:
        s=pd.read_csv(ROOT/f'data/derived/{c}_scores.tsv',sep='\t')
        require(s.panel_mean.notna().equals(s[GENES].notna().all(axis=1)), 'Integrity check failed: s.panel_mean.notna().equals(s[GENES].notna().all(axis=1))')
        v=s[s.panel_mean.notna()];np.testing.assert_allclose(v.panel_mean,v[GENES].mean(axis=1),rtol=1e-12)
    expr=pd.read_csv(ROOT/'results/expression_associations.tsv',sep='\t')
    require(len(expr) == 20 and expr.groupby('cohort').size().eq(10).all(), "Integrity check failed: len(expr) == 20 and expr.groupby('cohort').size().eq(10).all()")
    for _,g in expr.groupby('cohort'):np.testing.assert_allclose(g.q_BH,bh(g.p),equal_nan=True,rtol=1e-12)
    ctx=pd.read_csv(ROOT/'results/molecular_context.tsv',sep='\t')
    for _,g in ctx.groupby('analysis'):
        np.testing.assert_allclose(g.q_BH,bh(g.p),equal_nan=True,rtol=1e-12)
        np.testing.assert_allclose(g.adjusted_q_BH,bh(g.adjusted_p),equal_nan=True,rtol=1e-12)
    text=(ROOT/'manuscript/manuscript.md').read_text()
    for token in ['outputs are pending','result files are still pending','Writer note:','nan=']:
        require(token not in text, token)
    require('## References' in text, "Integrity check failed: '## References' in text")
    return dict(contrasts=len(d),estimated=len(estimated),expression_tests=len(expr),bootstrap=N_BOOT)


def check_workbooks():
    records=[]
    for p in list((ROOT/'submission').glob('*.xlsx'))+list((ROOT/'figures/source_data').glob('F[189]*.xlsx')):
        wb=load_workbook(p,read_only=True,data_only=False)
        sheets=[]
        for ws in wb.worksheets:
            errors=[c.coordinate for row in ws for c in row if c.data_type=='e']
            require(not errors, f'{p.name}/{ws.title}: {errors}')
            sheets.append(dict(sheet=ws.title,rows=ws.max_row,columns=ws.max_column))
        wb.close();records.append(dict(file=str(p.relative_to(ROOT)),sheets=sheets))
    return records


def structural_figures():
    validator=Path('tools/validate_figure_package.py')
    records=[]
    discovery=json.loads((ROOT/'registry/original_gene_selection.json').read_text())
    for source in discovery['sources']:
        require(sha256(ROOT / source['path']) == source['sha256'], source['path'])
    manifest=ROOT/'figures/F1_data_flow_manifest.json'
    flow=json.loads(manifest.read_text())
    for record in [flow['script']]+flow['inputs']+flow['outputs']+flow['source_data']:
        require(sha256(ROOT / record['path']) == record['sha256'], record['path'])
    status='not_run_installed_validator_unavailable'
    if validator.is_file():
        proc=subprocess.run([sys.executable,str(validator),'--base',str(ROOT),'--manifest',str(manifest)],
                            cwd=ROOT,capture_output=True,text=True)
        (ROOT/'verification/F1_data_flow_structure.log').write_text(proc.stdout+proc.stderr)
        if proc.returncode:raise AssertionError(f'Figure structural validation: F1_data_flow: {proc.stderr}')
        status='passed'
    records.append(dict(figure='F1_data_flow',structural_validation=status,
        provenance='Historical selection criteria and counts from preserved original sources; screen not recomputed',
        files=flow['outputs'],source_data=flow['source_data']))
    for stem in STEMS:
        paths=[ROOT/f'figures/{stem}.{ext}' for ext in ['png','pdf','svg']]
        for p in paths: require(p.is_file() and p.stat().st_size>1000, f'Missing or too-small evidence file: {p}')
        source=ROOT/f'figures/source_data/{stem}.xlsx'
        rec=dict(figure=stem,files=[dict(path=str(p.relative_to(ROOT)),sha256=sha256(p)) for p in paths],
                 source_data=dict(path=str(source.relative_to(ROOT)),sha256=sha256(source)))
        if validator.is_file():
            manifest=ROOT/f'verification/{stem}_manifest.json'
            cmd=[sys.executable,str(validator),'--base',str(ROOT),'--manifest',str(manifest),
                '--profile','publication','--figure-id',stem,'--script','scripts/make_figures.py',
                '--source-data',str(source),'--input','data/derived/all_public_analysis.tsv',
                '--input','results/tissue_contrasts.tsv','--input','results/meta_analysis.tsv',
                '--input','results/molecular_context.tsv','--input','results/expression_associations.tsv',
                '--claim','Fixed gene tissue and tumor-context associations shown in source tables',
                '--transformation','Beta scaled by100 for display; paired or Welch/cluster contrasts with patient bootstrap; gene order fixed',
                '--limitation','Gene-level arrays do not validate exact PSQ assay; cross-sectional comparisons are not progression; structural validation does not assess scientific claims',
                '--write-manifest']
            for p in paths:cmd+=['--figure',str(p)]
            proc=subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True)
            (ROOT/f'verification/{stem}_structure.log').write_text(proc.stdout+proc.stderr)
            if proc.returncode:raise AssertionError(f'Figure structural validation: {stem}: {proc.stderr}')
            rec['structural_validation']='passed'
        else:rec['structural_validation']='not_run_installed_validator_unavailable'
        records.append(rec)
    return records


def delivery_archive():
    entries=[]
    for p in sorted((ROOT/'submission').rglob('*')):
        if p.is_file() and p.name!='SHA256SUMS.txt':entries.append((p,str(p.relative_to(ROOT/'submission')),sha256(p)))
    (ROOT/'submission/SHA256SUMS.txt').write_text(''.join(f'{digest}  {name}\n' for _,name,digest in entries))
    archive=ROOT/'Public_Biology_Review_20260905.zip'
    with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted((ROOT/'submission').rglob('*')):
            if p.is_file():z.write(p,p.relative_to(ROOT/'submission'))
    with zipfile.ZipFile(archive) as z:
        require(z.testzip() is None, 'Integrity check failed: z.testzip() is None')
        import hashlib
        for _,name,digest in entries: require(hashlib.sha256(z.read(name)).hexdigest()==digest, f'ZIP member hash mismatch: {name}')
    write_json(ROOT/'verification/delivery_manifest.json',dict(archive=archive.name,sha256=sha256(archive),
        files=[dict(path=name,sha256=digest,bytes=p.stat().st_size) for p,name,digest in entries],external_upload=False))
    return dict(file=archive.name,sha256=sha256(archive),members=len(entries)+1)


def main():
    baseline=check_baseline();sources=check_downloads();results=check_results()
    books=check_workbooks();figures=structural_figures()
    render=json.loads((ROOT/'verification/render_manifest.json').read_text())
    for doc in render['documents']:
        require(sha256(ROOT / doc['document']) == doc['docx_sha256'], "Integrity check failed: sha256(ROOT / doc['document']) == doc['docx_sha256']")
        require(sha256(ROOT / doc['pdf']) == doc['pdf_sha256'], "Integrity check failed: sha256(ROOT / doc['pdf']) == doc['pdf_sha256']")
        require(not any((p['text_outside_page'] for p in doc['pages'])), "Integrity check failed: not any((p['text_outside_page'] for p in doc['pages']))")
    archive=delivery_archive()
    visual_path=ROOT/'verification/visual_review.json'
    current_visual=False
    if visual_path.exists():
        visual=json.loads(visual_path.read_text())
        current_visual=all((ROOT/f['path']).is_file() and sha256(ROOT/f['path'])==f['sha256'] for f in visual['files'])
    write_json(ROOT/'verification/package_verification.json',dict(status='numerical_and_structural_pass',
        baseline=baseline,downloaded_sources=sources,results=results,workbooks=books,figures=figures,
        rendered_documents=len(render['documents']),visual_review_current=current_visual,archive=archive))
    print(json.dumps(dict(status='numerical_and_structural_pass',results=results,
        baseline_files_per_copy=baseline['files_per_copy'],visual_review_current=current_visual,archive=archive),indent=2))


if __name__=='__main__':main()
